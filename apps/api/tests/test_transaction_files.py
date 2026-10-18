from __future__ import annotations

import csv
import uuid
from io import BytesIO, StringIO

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.infrastructure.database.models import ExportArtifactModel
from tests.helpers import auth_headers, register_and_get_token


async def _account(client: AsyncClient, token: str, name: str, balance: int = 100_000) -> dict:
    response = await client.post(
        '/api/v1/accounts',
        json={'name': name, 'type': 'bank', 'balance_paise': balance, 'currency': 'INR'},
        headers=auth_headers(token),
    )
    assert response.status_code == 201, response.text
    return response.json()


async def _category(client: AsyncClient, token: str, name: str, type: str) -> dict:
    response = await client.post(
        '/api/v1/categories', json={'name': name, 'type': type}, headers=auth_headers(token)
    )
    assert response.status_code == 201, response.text
    return response.json()


def _csv_bytes(rows: list[list[str]]) -> bytes:
    output = StringIO(newline='')
    writer = csv.writer(output)
    writer.writerows(rows)
    return output.getvalue().encode('utf-8')


async def test_export_formats_are_downloadable_and_owner_scoped(client: AsyncClient) -> None:
    owner = await register_and_get_token(client, 'report-owner@example.com')
    other = await register_and_get_token(client, 'report-other@example.com')
    account = await _account(client, owner, 'Axis')
    category = await _category(client, owner, 'Food', 'expense')
    created = await client.post(
        '/api/v1/transactions',
        json={
            'account_id': account['id'], 'category_id': category['id'], 'type': 'expense',
            'amount_paise': 12345, 'merchant': '=Cafe', 'note': 'Lunch',
            'occurred_at': '2026-10-15T12:00:00+05:30',
        },
        headers=auth_headers(owner),
    )
    assert created.status_code == 201, created.text

    for format in ('csv', 'xlsx', 'pdf'):
        response = await client.post(
            '/api/v1/exports/transactions',
            json={'format': format, 'date_from': '2026-10-15', 'date_to': '2026-10-15',
                  'timezone': 'Asia/Calcutta'},
            headers=auth_headers(owner),
        )
        assert response.status_code == 200, response.text
        payload = response.json()
        assert payload['row_count'] == 1
        assert payload['filename'].endswith('.' + format)
        assert (await client.get(
            f"/api/v1/exports/{payload['export_id']}/download-url", headers=auth_headers(other)
        )).status_code == 404
        file_response = await client.get(payload['download_url'])
        assert file_response.status_code == 200, file_response.text
        assert 'attachment' in file_response.headers['content-disposition']
        if format == 'csv':
            rows = list(csv.reader(StringIO(file_response.content.decode('utf-8-sig'))))
            assert rows[0] == [
                'occurred_at', 'type', 'amount_rupees', 'account', 'category', 'to_account', 'merchant',
            ]
            assert rows[1][2] == '123.45'
            assert rows[1][6] == "'=Cafe"
            assert account['id'] not in str(rows)
            assert category['id'] not in str(rows)
            assert created.json()['id'] not in str(rows)
            assert 'Lunch' not in str(rows)
        elif format == 'xlsx':
            workbook = load_workbook(BytesIO(file_response.content), read_only=True)
            try:
                rows = list(workbook.active.values)
                assert rows[0] == (
                    'occurred_at', 'type', 'amount_rupees', 'account', 'category', 'to_account', 'merchant',
                )
                assert rows[1][2] == 123.45
                assert rows[1][6] == "'=Cafe"
                assert account['id'] not in str(rows)
                assert category['id'] not in str(rows)
                assert created.json()['id'] not in str(rows)
                assert 'Lunch' not in str(rows)
            finally:
                workbook.close()
        else:
            assert file_response.content.startswith(b'%PDF')
            assert len(file_response.content) > 500
        assert (await client.get(f"/api/v1/exports/{payload['export_id']}/file?token=bad")).status_code == 404


async def test_old_report_with_internal_columns_cannot_be_downloaded(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    token = await register_and_get_token(client, 'report-old-format@example.com')
    response = await client.post(
        '/api/v1/exports/transactions', json={'format': 'csv'}, headers=auth_headers(token)
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    result = await db_session.execute(
        select(ExportArtifactModel).where(ExportArtifactModel.id == uuid.UUID(payload['export_id']))
    )
    artifact = result.scalar_one()
    artifact.report_schema_version = None
    await db_session.commit()
    assert (await client.get(payload['download_url'])).status_code == 410
    assert (await client.get(
        f"/api/v1/exports/{payload['export_id']}/download-url", headers=auth_headers(token)
    )).status_code == 410


async def test_csv_import_preview_confirm_balance_and_duplicate_protection(client: AsyncClient) -> None:
    token = await register_and_get_token(client, 'import-owner@example.com')
    axis = await _account(client, token, 'Axis')
    savings = await _account(client, token, 'Savings', balance=20_000)
    await _category(client, token, 'Food', 'expense')
    await _category(client, token, 'Salary', 'income')
    content = _csv_bytes([
        ['occurred_at', 'type', 'amount_paise', 'account', 'category', 'to_account', 'merchant'],
        ['2026-10-10T10:00:00+05:30', 'expense', '2500', 'Axis', 'Food', '', 'Cafe'],
        ['2026-10-11T10:00:00+05:30', 'income', '5000', 'Axis', 'Salary', '', 'Employer'],
        ['2026-10-12T10:00:00+05:30', 'transfer', '1000', 'Axis', '', 'Savings', ''],
    ])
    files = {'file': ('transactions.csv', content, 'text/csv')}
    preview = await client.post(
        '/api/v1/imports/transactions/preview', files=files, headers=auth_headers(token)
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()['row_count'] == preview.json()['valid_count'] == 3
    assert preview.json()['errors'] == []
    assert (await client.get('/api/v1/transactions', headers=auth_headers(token))).json() == []
    confirmed = await client.post(
        '/api/v1/imports/transactions/confirm', files=files,
        data={'preview_hash': preview.json()['file_hash']}, headers=auth_headers(token),
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()['imported_count'] == 3
    assert len((await client.get('/api/v1/transactions', headers=auth_headers(token))).json()) == 3
    axis_result = await client.get(f"/api/v1/accounts/{axis['id']}", headers=auth_headers(token))
    savings_result = await client.get(f"/api/v1/accounts/{savings['id']}", headers=auth_headers(token))
    assert axis_result.json()['balance_paise'] == 101_500
    assert savings_result.json()['balance_paise'] == 21_000
    repeat = await client.post(
        '/api/v1/imports/transactions/confirm', files=files,
        data={'preview_hash': preview.json()['file_hash']}, headers=auth_headers(token),
    )
    assert repeat.status_code == 409
    assert len((await client.get('/api/v1/transactions', headers=auth_headers(token))).json()) == 3


async def test_import_invalid_rows_or_foreign_account_make_no_writes(client: AsyncClient) -> None:
    owner = await register_and_get_token(client, 'import-private-owner@example.com')
    other = await register_and_get_token(client, 'import-private-other@example.com')
    foreign = await _account(client, other, 'Private')
    await _account(client, owner, 'Axis')
    await _category(client, owner, 'Food', 'expense')
    content = _csv_bytes([
        ['occurred_at', 'type', 'amount_paise', 'account_id', 'category'],
        ['2026-10-10', 'expense', '200', foreign['id'], 'Food'],
        ['2026-10-11', 'expense', '300', 'bad', 'Food'],
    ])
    files = {'file': ('invalid.csv', content, 'text/csv')}
    preview = await client.post('/api/v1/imports/transactions/preview', files=files, headers=auth_headers(owner))
    assert preview.status_code == 200
    assert len(preview.json()['errors']) == 2
    confirm = await client.post(
        '/api/v1/imports/transactions/confirm', files=files,
        data={'preview_hash': preview.json()['file_hash']}, headers=auth_headers(owner),
    )
    assert confirm.status_code == 400
    assert (await client.get('/api/v1/transactions', headers=auth_headers(owner))).json() == []


async def test_xlsx_import_uses_owned_account_and_category(client: AsyncClient) -> None:
    token = await register_and_get_token(client, 'import-xlsx@example.com')
    await _account(client, token, 'Axis')
    await _category(client, token, 'Food', 'expense')
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['date', 'type', 'amount', 'account', 'category', 'merchant'])
    sheet.append(['2026-10-12', 'expense', '12.34', 'Axis', 'Food', 'Bakery'])
    output = BytesIO()
    workbook.save(output)
    files = {'file': ('transactions.xlsx', output.getvalue(),
                      'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    preview = await client.post('/api/v1/imports/transactions/preview', files=files, headers=auth_headers(token))
    assert preview.status_code == 200, preview.text
    assert preview.json()['valid_count'] == 1
    assert preview.json()['preview_rows'][0]['amount_paise'] == 1234
    confirmed = await client.post(
        '/api/v1/imports/transactions/confirm', files=files,
        data={'preview_hash': preview.json()['file_hash']}, headers=auth_headers(token),
    )
    assert confirmed.status_code == 200, confirmed.text


async def test_csv_export_import_maps_names_to_another_users_owned_records(client: AsyncClient) -> None:
    source = await register_and_get_token(client, 'report-move-source@example.com')
    target = await register_and_get_token(client, 'report-move-target@example.com')
    source_account = await _account(client, source, 'Axis')
    source_category = await _category(client, source, 'Food', 'expense')
    await _account(client, target, 'Axis')
    await _category(client, target, 'Food', 'expense')
    created = await client.post(
        '/api/v1/transactions',
        json={
            'account_id': source_account['id'], 'category_id': source_category['id'],
            'type': 'expense', 'amount_paise': 1234, 'merchant': 'Cafe',
            'occurred_at': '2026-10-15T12:00:00+05:30',
        },
        headers=auth_headers(source),
    )
    assert created.status_code == 201, created.text
    exported = await client.post(
        '/api/v1/exports/transactions', json={'format': 'csv'}, headers=auth_headers(source)
    )
    assert exported.status_code == 200, exported.text
    download = await client.get(exported.json()['download_url'])
    files = {'file': ('transactions.csv', download.content, 'text/csv')}
    source_preview = await client.post(
        '/api/v1/imports/transactions/preview', files=files, headers=auth_headers(source)
    )
    assert source_preview.status_code == 200, source_preview.text
    assert source_preview.json()['already_imported'] is True
    source_confirm = await client.post(
        '/api/v1/imports/transactions/confirm', files=files,
        data={'preview_hash': source_preview.json()['file_hash']}, headers=auth_headers(source),
    )
    assert source_confirm.status_code == 409
    preview = await client.post('/api/v1/imports/transactions/preview', files=files, headers=auth_headers(target))
    assert preview.status_code == 200, preview.text
    assert preview.json()['already_imported'] is False
    assert preview.json()['valid_count'] == 1
    confirmed = await client.post(
        '/api/v1/imports/transactions/confirm', files=files,
        data={'preview_hash': preview.json()['file_hash']}, headers=auth_headers(target),
    )
    assert confirmed.status_code == 200, confirmed.text
    imported = (await client.get('/api/v1/transactions', headers=auth_headers(target))).json()
    assert len(imported) == 1
    assert imported[0]['amount_paise'] == 1234
    assert imported[0]['merchant'] == 'Cafe'


async def test_import_later_transfer_failure_rolls_back_earlier_rows(client: AsyncClient) -> None:
    token = await register_and_get_token(client, 'import-atomic@example.com')
    axis = await _account(client, token, 'Axis', balance=1_000)
    savings = await _account(client, token, 'Savings', balance=0)
    await _category(client, token, 'Salary', 'income')
    content = _csv_bytes([
        ['occurred_at', 'type', 'amount_paise', 'account', 'category', 'to_account'],
        ['2026-10-12', 'transfer', '2000', 'Axis', '', 'Savings'],
        ['2026-10-10', 'income', '100', 'Axis', 'Salary', ''],
    ])
    files = {'file': ('atomic.csv', content, 'text/csv')}
    preview = await client.post('/api/v1/imports/transactions/preview', files=files, headers=auth_headers(token))
    assert preview.status_code == 200, preview.text
    assert preview.json()['valid_count'] == 2
    confirmed = await client.post(
        '/api/v1/imports/transactions/confirm', files=files,
        data={'preview_hash': preview.json()['file_hash']}, headers=auth_headers(token),
    )
    assert confirmed.status_code == 400
    assert (await client.get('/api/v1/transactions', headers=auth_headers(token))).json() == []
    axis_after = await client.get(f"/api/v1/accounts/{axis['id']}", headers=auth_headers(token))
    savings_after = await client.get(f"/api/v1/accounts/{savings['id']}", headers=auth_headers(token))
    assert axis_after.json()['balance_paise'] == 1000
    assert savings_after.json()['balance_paise'] == 0

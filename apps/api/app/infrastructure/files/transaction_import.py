from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from itertools import islice
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.domain.exceptions import InvalidImportError

MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 10_000
MAX_UNCOMPRESSED_XLSX_BYTES = 50 * 1024 * 1024
MAX_AMOUNT_PAISE = 2**63 - 1


@dataclass(frozen=True)
class RawImportRow:
    row_number: int
    values: dict[str, str]


def _header(value: Any) -> str:
    text = str(value or '').strip().casefold().replace(' ', '_').replace('-', '_')
    return {
        'date': 'occurred_at',
        'datetime': 'occurred_at',
        'timestamp': 'occurred_at',
        'transaction_date': 'occurred_at',
        'amount': 'amount_rupees',
        'account_name': 'account',
        'category_name': 'category',
        'destination_account': 'to_account',
        'description': 'note',
    }.get(text, text)


def _cell(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    if text.startswith("'") and text[1:2] in {'=', '+', '-', '@'}:
        return text[1:]
    return text


def parse_transaction_file(filename: str, content: bytes) -> list[RawImportRow]:
    extension = Path(filename).suffix.casefold()
    if extension not in {'.csv', '.xlsx'}:
        raise InvalidImportError('Choose a CSV or XLSX transaction file.')
    if not content or len(content) > MAX_IMPORT_BYTES:
        raise InvalidImportError('The import file must be nonempty and no larger than 10 MB.')
    try:
        if extension == '.csv':
            reader = csv.reader(StringIO(content.decode('utf-8-sig'), newline=''))
            records = list(islice(reader, MAX_IMPORT_ROWS + 2))
        else:
            with ZipFile(BytesIO(content)) as archive:
                if len(archive.infolist()) > 1_000 or sum(
                    entry.file_size for entry in archive.infolist()
                ) > MAX_UNCOMPRESSED_XLSX_BYTES:
                    raise InvalidImportError('The XLSX file is too large to import.')
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                records = list(islice(workbook.active.iter_rows(values_only=True), MAX_IMPORT_ROWS + 2))
            finally:
                workbook.close()
    except (UnicodeError, csv.Error, OSError, ValueError, KeyError, BadZipFile) as exc:
        raise InvalidImportError('Could not read this CSV or XLSX file.') from exc
    if len(records) < 2:
        raise InvalidImportError('The file needs a header and at least one transaction.')
    if len(records) - 1 > MAX_IMPORT_ROWS:
        raise InvalidImportError('Import at most 10,000 transactions at once.')
    headers = [_header(value) for value in records[0]]
    if not headers or any(not header for header in headers) or len(set(headers)) != len(headers):
        raise InvalidImportError('The file has missing or duplicate column headings.')
    required = {'occurred_at', 'type'}
    if not required.issubset(headers) or not ({'amount_paise', 'amount_rupees'} & set(headers)):
        raise InvalidImportError('The file needs occurred_at, type, and amount_paise or amount columns.')
    if not ({'account', 'account_id'} & set(headers)):
        raise InvalidImportError('The file needs an account or account_id column.')
    rows = []
    for row_number, record in enumerate(records[1:], start=2):
        if all(not _cell(value) for value in record):
            continue
        if len(record) > len(headers) and any(_cell(value) for value in record[len(headers):]):
            raise InvalidImportError(f'Row {row_number} has more cells than the header.')
        rows.append(
            RawImportRow(
                row_number=row_number,
                values={header: _cell(record[index]) if index < len(record) else ''
                        for index, header in enumerate(headers)},
            )
        )
    if not rows:
        raise InvalidImportError('The file has no transactions to import.')
    return rows


def parse_amount_paise(values: dict[str, str]) -> int:
    raw_paise = values.get('amount_paise', '').replace(',', '')
    if raw_paise:
        if len(raw_paise) > 19 or not raw_paise.isascii() or not raw_paise.isdecimal():
            raise ValueError('amount_paise must be a whole number')
        try:
            amount = int(raw_paise)
        except ValueError as exc:
            raise ValueError('amount_paise must be a whole number') from exc
    else:
        raw_rupees = values.get('amount_rupees', '').replace(',', '').replace('₹', '').strip()
        if len(raw_rupees) > 32 or not re.fullmatch(r'[0-9]+(?:\.[0-9]{1,2})?', raw_rupees):
            raise ValueError('amount must be a positive number with at most two decimal places')
        try:
            amount_rupees = Decimal(raw_rupees)
        except InvalidOperation as exc:
            raise ValueError('amount must be a number') from exc
        if not amount_rupees.is_finite() or amount_rupees * 100 != (amount_rupees * 100).to_integral_value():
            raise ValueError('amount must have at most two decimal places')
        amount = int(amount_rupees * 100)
    if amount <= 0 or amount > MAX_AMOUNT_PAISE:
        raise ValueError('amount must be greater than zero and fit the supported range')
    return amount


def parse_occurred_at(raw: str, timezone: ZoneInfo) -> datetime:
    try:
        occurred_at = datetime.fromisoformat(raw.replace('Z', '+00:00'))
    except ValueError as exc:
        raise ValueError('occurred_at must be an ISO date or datetime') from exc
    if occurred_at.tzinfo is None:
        occurred_at = occurred_at.replace(tzinfo=timezone)
    return occurred_at
